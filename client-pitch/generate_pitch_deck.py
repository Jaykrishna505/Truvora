from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent / "Truvora_Client_Pitch_Deck.pptx"

WIDE_W = Inches(13.333)
WIDE_H = Inches(7.5)

INK = RGBColor(23, 33, 31)
MUTED = RGBColor(91, 108, 102)
GREEN = RGBColor(22, 122, 91)
GREEN_DARK = RGBColor(12, 70, 54)
GREEN_SOFT = RGBColor(229, 244, 237)
GOLD = RGBColor(155, 108, 24)
GOLD_SOFT = RGBColor(251, 240, 213)
BLUE = RGBColor(49, 95, 159)
WHITE = RGBColor(255, 255, 255)
PAGE = RGBColor(247, 250, 248)
LINE = RGBColor(220, 229, 224)
RED = RGBColor(189, 79, 61)


def add_textbox(slide, x, y, w, h, text, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_multiline(slide, x, y, w, h, lines, size=17, color=MUTED, bullet=False, line_spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    for idx, line in enumerate(lines):
        para = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        para.text = line
        para.font.name = "Aptos"
        para.font.size = Pt(size)
        para.font.color.rgb = color
        para.level = 0
        para.line_spacing = line_spacing
        if bullet:
            para.text = f"• {line}"
    return box


def add_rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    shape.line.width = Pt(1)
    return shape


def add_title(slide, eyebrow, title, subtitle=None):
    add_textbox(slide, 0.7, 0.48, 3.8, 0.25, eyebrow.upper(), 10, GREEN, True)
    add_textbox(slide, 0.7, 0.82, 8.8, 0.9, title, 33, INK, True)
    if subtitle:
        add_textbox(slide, 0.72, 1.72, 8.9, 0.45, subtitle, 16, MUTED)


def add_footer(slide, cue):
    add_rect(slide, 0, 7.13, 13.333, 0.37, GREEN_DARK)
    add_textbox(slide, 0.7, 7.21, 11.8, 0.18, f"Talk track: {cue}", 8.5, WHITE)


def add_background(slide, image_name):
    image = ROOT / "app" / "static" / image_name
    slide.shapes.add_picture(str(image), 0, 0, width=WIDE_W, height=WIDE_H)
    add_rect(slide, 0, 0, 13.333, 7.5, RGBColor(255, 255, 255))
    overlay = slide.shapes[-1]
    overlay.fill.transparency = 18
    overlay.line.transparency = 100


def add_card(slide, x, y, w, h, title, body, accent=GREEN):
    add_rect(slide, x, y, w, h, WHITE, LINE, radius=True)
    add_rect(slide, x, y, 0.08, h, accent)
    add_textbox(slide, x + 0.24, y + 0.22, w - 0.45, 0.35, title, 18, INK, True)
    add_multiline(slide, x + 0.24, y + 0.72, w - 0.45, h - 0.9, body, 13.5, MUTED, bullet=True, line_spacing=1.06)


def add_package_card(slide, x, package):
    y = 1.55
    w = 3.8
    h = 4.85
    accent = package["accent"]
    add_rect(slide, x, y, w, h, WHITE, LINE, radius=True)
    add_rect(slide, x, y, w, 0.12, accent)
    add_textbox(slide, x + 0.25, y + 0.32, w - 0.5, 0.32, package["name"], 18, INK, True)
    add_textbox(slide, x + 0.25, y + 0.82, 1.7, 0.45, f"${package['price']}", 32, accent, True)
    add_textbox(slide, x + 1.85, y + 1.04, 1.4, 0.25, "/ month", 12, MUTED, True)
    add_textbox(slide, x + 0.25, y + 1.42, w - 0.5, 0.6, package["best"], 13, MUTED)
    add_multiline(slide, x + 0.28, y + 2.18, w - 0.55, 1.55, package["includes"], 11.7, INK, bullet=True)
    add_rect(slide, x + 0.25, y + 4.06, w - 0.5, 0.52, package["pill"], package["pill"], radius=True)
    add_textbox(slide, x + 0.42, y + 4.2, w - 0.84, 0.2, package["position"], 10.5, accent, True, PP_ALIGN.CENTER)


def add_flow_step(slide, idx, x, title, detail, accent=GREEN):
    add_rect(slide, x, 2.08, 1.45, 1.45, accent, accent, radius=True)
    add_textbox(slide, x + 0.45, 2.43, 0.55, 0.38, str(idx), 26, WHITE, True, PP_ALIGN.CENTER)
    add_textbox(slide, x - 0.08, 3.78, 1.65, 0.36, title, 14, INK, True, PP_ALIGN.CENTER)
    add_textbox(slide, x - 0.25, 4.22, 2.0, 0.54, detail, 10.5, MUTED, False, PP_ALIGN.CENTER)


prs = Presentation()
prs.slide_width = WIDE_W
prs.slide_height = WIDE_H
blank = prs.slide_layouts[6]


# 1. Cover
slide = prs.slides.add_slide(blank)
add_background(slide, "auth-background.jpg")
add_rect(slide, 0, 0, 13.333, 7.5, RGBColor(8, 34, 28))
slide.shapes[-1].fill.transparency = 28
add_rect(slide, 0.72, 0.58, 0.48, 0.48, GREEN, GREEN, radius=True)
add_textbox(slide, 0.88, 0.68, 0.16, 0.18, "T", 16, WHITE, True, PP_ALIGN.CENTER)
add_textbox(slide, 1.32, 0.67, 2.0, 0.28, "Truvora", 18, WHITE, True)
add_textbox(slide, 0.72, 2.1, 8.35, 1.25, "Turn guest feedback into reviews, reputation, and growth.", 42, WHITE, True)
add_textbox(slide, 0.76, 3.55, 6.6, 0.58, "A simple platform and service package for hotels that want more review opportunities, better guest insight, and a stronger online presence.", 17, RGBColor(229, 244, 237))
add_rect(slide, 0.74, 4.55, 2.35, 0.54, GOLD_SOFT, GOLD_SOFT, radius=True)
add_textbox(slide, 0.92, 4.72, 1.98, 0.15, "Client Pitch Deck", 11, GOLD, True, PP_ALIGN.CENTER)
add_footer(slide, "Open with the business outcome: more trust online, less manual follow-up, and a system hotels can actually use.")

# 2. Problem
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAGE
add_title(slide, "The challenge", "Hotels lose reputation momentum when guest feedback is manual.", "Most hotels know reviews matter. The hard part is asking consistently, responding quickly, and organizing every signal.")
add_card(slide, 0.7, 2.45, 3.7, 2.6, "Manual follow-up breaks", ["Front desk teams get busy", "Guests leave before being asked", "No single place to track outreach"], GREEN)
add_card(slide, 4.8, 2.45, 3.7, 2.6, "Negative feedback gets missed", ["Issues stay hidden until public reviews", "Managers lack fast visibility", "Follow-up notes are scattered"], RED)
add_card(slide, 8.9, 2.45, 3.7, 2.6, "Online presence goes stale", ["Review replies fall behind", "Social pages look inactive", "Future guests see silence"], BLUE)
add_footer(slide, "Frame the pain gently: they are not doing anything wrong; they just need a repeatable process.")

# 3. Solution
slide = prs.slides.add_slide(blank)
add_background(slide, "dashboard-background.jpg")
add_rect(slide, 0.68, 0.62, 5.6, 5.82, WHITE, WHITE, radius=True)
slide.shapes[-1].fill.transparency = 4
add_textbox(slide, 1.05, 1.05, 1.25, 0.24, "SOLUTION", 10, GREEN, True)
add_textbox(slide, 1.04, 1.45, 4.45, 0.85, "One workflow for guest outreach and reputation follow-up.", 30, INK, True)
add_multiline(slide, 1.08, 2.68, 4.35, 1.75, ["Send SMS/email review requests", "Capture private feedback and comments", "Track sent, opened, responded, and Google clicks", "Keep notes and follow-ups in one dashboard"], 15, MUTED, True)
add_rect(slide, 1.05, 5.25, 3.9, 0.55, GREEN, GREEN, radius=True)
add_textbox(slide, 1.28, 5.42, 3.42, 0.16, "Built for hotel teams, not tech teams", 12, WHITE, True, PP_ALIGN.CENTER)
add_footer(slide, "Position this as operational simplicity: one place to request, record, and act.")

# 4. Flow
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = WHITE
add_title(slide, "Platform flow", "From guest stay to actionable feedback.", "The platform keeps the process simple for staff and easy for guests.")
steps = [
    ("Hotel setup", "Add Google review link and message templates."),
    ("Guest entry", "Enter guest name, phone, email, and stay date."),
    ("Request sent", "SMS/email sends a unique feedback link."),
    ("Guest rates", "Guest selects stars and may add comments."),
    ("Dashboard updates", "Owner sees feedback, status, and notes."),
]
for i, (title, detail) in enumerate(steps, 1):
    add_flow_step(slide, i, 0.95 + (i - 1) * 2.45, title, detail, GREEN if i < 5 else BLUE)
    if i < 5:
        add_textbox(slide, 2.45 + (i - 1) * 2.45, 2.6, 0.45, 0.24, ">", 24, LINE, True, PP_ALIGN.CENTER)
add_footer(slide, "Walk them through this as a real operational process, not just software features.")

# 5. Compliance
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAGE
add_title(slide, "Compliant review handling", "We record feedback without blocking public review access.", "Google review behavior has to be handled carefully. The platform is designed around that reality.")
add_card(slide, 0.85, 2.25, 3.6, 2.8, "What we can record", ["Star rating selected on our page", "Comments entered before redirect", "Whether Google review button was clicked"], GREEN)
add_card(slide, 4.86, 2.25, 3.6, 2.8, "What we cannot record", ["Whether a Google review was submitted", "What the guest posted on Google", "Any private Google account action"], GOLD)
add_card(slide, 8.87, 2.25, 3.6, 2.8, "Why this matters", ["Keeps process transparent", "Avoids unfair review gating", "Protects hotel reputation practices"], BLUE)
add_footer(slide, "This builds trust. Be clear: we track the intent and internal feedback, not private activity on Google.")

# 6. Packages
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = WHITE
add_title(slide, "Service packages", "Choose the level of support your hotel needs.", "Start with outreach, add reputation support, or include social media management.")
packages = [
    {
        "name": "Guest Outreach",
        "price": "149",
        "best": "Best for hotels that want consistent guest review requests.",
        "includes": ["SMS review requests", "Email review requests", "Feedback dashboard", "Status and comments"],
        "position": "More feedback opportunities",
        "accent": GREEN,
        "pill": GREEN_SOFT,
    },
    {
        "name": "Reputation Management",
        "price": "249",
        "best": "Best for hotels that want help replying to reviews.",
        "includes": ["Everything in Outreach", "Review reply workflow", "Professional response support", "Follow-up notes"],
        "position": "Better public reputation",
        "accent": BLUE,
        "pill": RGBColor(231, 238, 249),
    },
    {
        "name": "Social Management",
        "price": "349",
        "best": "Best for hotels that want reputation and visibility support.",
        "includes": ["Everything in Reputation", "Facebook post management", "Social content coordination", "Premium workflow"],
        "position": "Reviews plus visibility",
        "accent": GOLD,
        "pill": GOLD_SOFT,
    },
]
for x, package in zip([0.68, 4.76, 8.84], packages):
    add_package_card(slide, x, package)
add_footer(slide, "Do not sell price first. Sell the level of workload they want removed from their team.")

# 7. Guest Outreach
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAGE
add_title(slide, "Package 1", "Guest Outreach: $149/month", "For hotels that want a simple system to ask guests for reviews after their stay.")
add_card(slide, 0.78, 2.15, 5.75, 3.25, "Included", ["SMS and email review request sending", "Guest feedback page with star rating and comments", "Dashboard for sent, opened, responded status", "Google review link tracking and follow-up notes"], GREEN)
add_card(slide, 6.9, 2.15, 5.75, 3.25, "Client talk track", ["Helps staff stay consistent without manual reminders", "Gives owners visibility into guest sentiment", "Best starter option when the main goal is more review opportunities", "Great fit for teams already handling replies themselves"], BLUE)
add_footer(slide, "Use this package when the client says they mainly need more reviews or a better request process.")

# 8. Reputation
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = WHITE
add_title(slide, "Package 2", "Reputation Management: $249/month", "For hotels that want outreach plus support managing and replying to reviews.")
add_card(slide, 0.78, 2.15, 5.75, 3.25, "Included", ["Everything in Guest Outreach", "Online review reply management", "Professional response drafting", "Follow-up workflow for guest issues"], BLUE)
add_card(slide, 6.9, 2.15, 5.75, 3.25, "Client talk track", ["Public replies are also read by future guests", "Fast, thoughtful replies reduce damage from negative reviews", "Shows the hotel is active and attentive", "Useful when staff do not have time to manage every reply"], GREEN)
add_footer(slide, "This is the natural upgrade for owners who care about perception, not just review volume.")

# 9. Social
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PAGE
add_title(slide, "Package 3", "Social Management: $349/month", "For hotels that want guest outreach, review support, and social media visibility.")
add_card(slide, 0.78, 2.15, 5.75, 3.25, "Included", ["Everything in Reputation Management", "Facebook post management", "Social content coordination", "Visibility workflow for hotel updates", "Premium service support"], GOLD)
add_card(slide, 6.9, 2.15, 5.75, 3.25, "Client talk track", ["Reviews build trust while social keeps the hotel visible", "An active online presence supports future bookings", "Good option for owners who want fewer online tasks to manage", "Useful for properties with events, offers, or seasonal updates"], GREEN)
add_footer(slide, "Sell this as the full online presence package: feedback, reputation, and visibility together.")

# 10. Dashboard
slide = prs.slides.add_slide(blank)
add_background(slide, "guest-background.jpg")
add_rect(slide, 0.68, 0.65, 11.98, 5.85, WHITE, WHITE, radius=True)
slide.shapes[-1].fill.transparency = 5
add_title(slide, "What the owner sees", "A dashboard built around action.", None)
add_card(slide, 0.95, 2.05, 3.65, 2.75, "Performance", ["Total requests sent", "Positive and negative feedback", "Open and response activity"], GREEN)
add_card(slide, 4.85, 2.05, 3.65, 2.75, "Guest records", ["Star rating and comments", "Google click status", "Stay date and contact details"], BLUE)
add_card(slide, 8.75, 2.05, 3.65, 2.75, "Follow-up", ["Internal notes", "Guest issue tracking", "Real-time updates"], GOLD)
add_footer(slide, "Show that the dashboard is not just reporting; it helps the owner decide who needs follow-up.")

# 11. Trial and billing
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = WHITE
add_title(slide, "Trial and billing", "Simple monthly packages with clear service rules.", "The platform is designed to let hotels try it, then continue only when they choose a paid package.")
add_card(slide, 0.85, 2.1, 3.65, 3.05, "7-day trial", ["Every hotel account starts with a 7-day trial", "Trial lets the owner test the workflow", "Paid package required after trial"], GREEN)
add_card(slide, 4.85, 2.1, 3.65, 3.05, "Monthly renewal", ["Packages renew monthly", "Package changes apply next renewal", "Current service remains active until then"], BLUE)
add_card(slide, 8.85, 2.1, 3.65, 3.05, "Cancellation", ["60-day notice to cancel", "Service stays active during notice", "Dashboard shows expiry date"], GOLD)
add_footer(slide, "Keep this simple. The client should feel the billing terms are clear and predictable.")

# 12. Close
slide = prs.slides.add_slide(blank)
add_background(slide, "auth-background.jpg")
add_rect(slide, 0, 0, 13.333, 7.5, RGBColor(8, 34, 28))
slide.shapes[-1].fill.transparency = 20
add_textbox(slide, 0.82, 0.75, 1.25, 0.24, "NEXT STEP", 10, GOLD_SOFT, True)
add_textbox(slide, 0.82, 1.28, 7.4, 1.15, "Start with a 7-day trial and choose the package that matches your goals.", 38, WHITE, True)
add_multiline(slide, 0.86, 2.86, 6.3, 1.25, ["Guest Outreach if you need more review opportunities", "Reputation Management if you need review reply support", "Social Management if you want reputation plus visibility"], 16, RGBColor(229, 244, 237), True)
add_rect(slide, 0.86, 4.75, 3.95, 0.66, GREEN, GREEN, radius=True)
add_textbox(slide, 1.15, 4.96, 3.35, 0.16, "Let us set up your hotel account", 13, WHITE, True, PP_ALIGN.CENTER)
add_footer(slide, "Close by asking which outcome matters most: more reviews, better reputation, or stronger online presence.")

prs.save(OUT)
print(OUT)
